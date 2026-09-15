"""
Тесты выгрузки меню в Excel.

Книгу нельзя проверить «на глаз в тесте», поэтому проверяется то, что
ломается молча: сходятся ли числа в таблице с самим меню, не потерялись ли
строки, и — главное — не обрезал ли Excel шкалу графика. Последнее особенно
коварно: обрезанная ось превращает отклонение в один процент в обвал,
и заметить это можно только глазами, если специально не проверить.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from export import menu_to_workbook, menu_to_xlsx
from solver import build_menu
from targets import Profile, compute_targets

PROFILE = Profile(sex="female", age=31, height_cm=168, weight_kg=74,
                  activity="light", goal="lose")


@pytest.fixture(scope="module")
def menu(catalog):
    return build_menu(catalog, compute_targets(PROFILE), days=5, seed=0)


@pytest.fixture(scope="module")
def book(menu):
    """Книга, прошедшая через сохранение и чтение.

    Именно так её увидит пользователь: openpyxl многое проверяет только
    при записи, и объект в памяти может быть исправным, а файл — нет.
    """
    return load_workbook(io.BytesIO(menu_to_xlsx(menu, PROFILE.model_dump())))


class TestStructure:
    def test_three_sheets_in_reading_order(self, book):
        assert book.sheetnames == ["Сводка", "По дням", "Меню"]

    def test_every_meal_became_a_row(self, book, menu):
        sheet = book["Меню"]
        expected = sum(len(day.items) for day in menu.days)
        assert sheet.max_row - 1 == expected   # минус строка заголовка

    def test_day_column_covers_all_days(self, book, menu):
        sheet = book["Меню"]
        days = {sheet.cell(row=r, column=1).value for r in range(2, sheet.max_row + 1)}
        assert days == set(range(1, len(menu.days) + 1))

    def test_menu_sheet_is_filterable(self, book):
        """Без автофильтра и закреплённой шапки таблица на 200 строк бесполезна."""
        sheet = book["Меню"]
        assert sheet.auto_filter.ref is not None
        assert sheet.freeze_panes == "A2"


class TestNumbers:
    def test_daily_totals_match_the_menu(self, book, menu):
        sheet = book["По дням"]
        for number, day in enumerate(menu.days, start=1):
            assert sheet.cell(row=number + 1, column=2).value == round(day.totals["kcal"])

    def test_target_column_is_the_norm(self, book, menu):
        sheet = book["По дням"]
        column = {sheet.cell(row=r, column=3).value for r in range(2, sheet.max_row + 1)}
        assert column == {round(menu.targets["kcal"])}

    def test_summary_average_matches_days(self, book, menu):
        """Среднесуточные калории на «Сводке» — это среднее по листу «По дням»."""
        expected = sum(day.totals["kcal"] for day in menu.days) / len(menu.days)
        assert book["Сводка"].cell(row=5, column=3).value == round(expected)

    def test_meal_rows_sum_to_daily_totals(self, book, menu):
        """Лист «Меню» и лист «По дням» не должны разъезжаться."""
        sheet = book["Меню"]
        first_day = [
            sheet.cell(row=r, column=5).value
            for r in range(2, sheet.max_row + 1)
            if sheet.cell(row=r, column=1).value == 1
        ]
        # Округление до целых по каждой строке даёт погрешность в пару ккал.
        assert abs(sum(first_day) - menu.days[0].totals["kcal"]) < len(first_day)


class TestCharts:
    def test_summary_has_all_three_charts(self, book):
        assert len(book["Сводка"]._charts) == 3

    def test_axes_start_from_a_fixed_floor(self, book):
        """Excel по умолчанию обрезает шкалу под данные — и врёт масштабом.

        На ровном рационе разброс в полтора процента растягивается на всю
        высоту графика, и меню, попадающее в норму, выглядит как хаос.
        """
        bars, line = book["Сводка"]._charts[1], book["Сводка"]._charts[2]
        assert bars.y_axis.scaling.min == 0
        assert line.y_axis.scaling.min is not None
        assert line.y_axis.scaling.max is not None

    def test_line_chart_shows_the_norm_alongside_the_fact(self, book):
        """Без линии нормы график калорий не с чем сравнить."""
        assert len(book["Сводка"]._charts[2].series) == 2

    def test_pie_slices_keep_the_fixed_macro_colours(self, book):
        """Белки синие, жиры оранжевые, углеводы бирюзовые — на всей книге."""
        from export import CARB, FAT, PROTEIN

        points = book["Сводка"]._charts[0].series[0].data_points
        assert [p.graphicalProperties.solidFill.srgbClr for p in points] == \
            [PROTEIN, FAT, CARB]

    def test_labels_do_not_repeat_the_series_name(self, book):
        """Иначе Excel пишет на секторе «Калорий в день, Белки, 35%»."""
        labels = book["Сводка"]._charts[0].dataLabels
        assert labels.showSerName is False
        assert labels.showPercent is True


class TestEdgeCases:
    def test_single_day_menu_still_builds(self, catalog):
        menu = build_menu(catalog, compute_targets(PROFILE), days=1, seed=0)
        book = load_workbook(io.BytesIO(menu_to_xlsx(menu)))
        assert book["По дням"].max_row == 2

    def test_profile_is_optional(self, menu):
        """Выгрузка не должна падать, если о человеке ничего не известно."""
        assert menu_to_workbook(menu, None) is not None

    def test_age_agrees_with_russian_grammar(self):
        from export import _years

        assert [_years(n) for n in (1, 3, 5, 11, 21, 22, 45)] == \
            ["1 год", "3 года", "5 лет", "11 лет", "21 год", "22 года", "45 лет"]
